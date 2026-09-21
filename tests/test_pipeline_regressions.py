import copy
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
import detect_capabilities as capabilities
import legal_common
import md2docx_legal as converter
import verify_docx as verifier


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.params = copy.deepcopy(legal_common.DEFAULT_PARAMS)

    def test_provider_slot_and_token_boundaries(self):
        with tempfile.TemporaryDirectory() as directory:
            profile = Path(directory) / 'profile.md'
            profile.write_text('LAW.*: yuandian pandoc\nKB.*: image\n', encoding='utf-8')
            declared = capabilities.parse_profile(profile)
            self.assertEqual(declared['LAW'], ['yuandian'])
            self.assertEqual(declared['KB'], [])
            profile.write_text('KB.*: IMA\nLAW.*: 元典\n', encoding='utf-8')
            declared = capabilities.parse_profile(profile)
            self.assertEqual(declared['KB'], ['ima'])
            self.assertEqual(declared['LAW'], ['yuandian'])

    def test_page_fix_updates_original_document(self):
        doc = Document()
        doc.sections[0].page_width = Inches(5)
        self.assertFalse(verifier.verify_v1(doc, self.params)['passed'])
        converter._auto_fix(doc, [{'check': 'V1'}, {'check': 'V2'}], self.params)
        self.assertTrue(verifier.verify_v1(doc, self.params)['passed'])
        self.assertTrue(verifier.verify_v2(doc, self.params)['passed'])

    def test_missing_signature_alignment_is_rejected(self):
        doc = Document()
        paragraph = doc.add_paragraph('具状人：测试')
        self.assertFalse(converter._verify_v10_signature_align(doc, self.params)['passed'])
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self.assertTrue(converter._verify_v10_signature_align(doc, self.params)['passed'])

    def test_configurable_alignment(self):
        doc = Document()
        title = doc.add_paragraph('测试标题')
        signature = doc.add_paragraph('具状人：测试')
        self.params['title']['align'] = 'left'
        self.params['signature']['align'] = 'left'
        self.assertTrue(verifier.verify_v4(doc, self.params)['passed'])
        self.assertTrue(verifier.verify_v10(doc, self.params)['passed'])
        self.params['title']['align'] = 'right'
        title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self.params['signature']['align'] = 'center'
        signature.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.assertTrue(verifier.verify_v4(doc, self.params)['passed'])
        self.assertTrue(verifier.verify_v10(doc, self.params)['passed'])

    def test_pandoc_receives_existing_reference_template(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'input.md'
            source.write_text('# Test\n', encoding='utf-8')
            with patch.object(converter, 'DocxBuilder', side_effect=RuntimeError('force fallback')):
                with patch('subprocess.run', return_value=subprocess.CompletedProcess([], 0)) as run:
                    result = converter.convert(str(source), str(Path(directory) / 'output.docx'))
            self.assertEqual(result['tier'], 'Tier 3 (pandoc)')
            reference = ROOT / 'scripts' / 'templates' / 'reference.docx'
            self.assertTrue(reference.is_file())
            self.assertIn('--reference-doc=' + str(reference), run.call_args.args[0])

    def test_cli_help_under_gbk_redirection(self):
        for script in ('md2docx_legal.py', 'verify_docx.py'):
            with self.subTest(script=script):
                result = subprocess.run([sys.executable, str(ROOT / 'scripts' / script), '--help'],
                                        env={**os.environ, 'PYTHONIOENCODING': 'gbk'}, capture_output=True)
                self.assertEqual(result.returncode, 0, result.stderr)
                result.stdout.decode('utf-8')


if __name__ == '__main__':
    unittest.main()
